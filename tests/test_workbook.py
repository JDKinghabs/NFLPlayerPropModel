import pandas as pd
from openpyxl import load_workbook
from props.pipeline import Result
from props.workbook import write_workbook


def td_rows():
    return pd.DataFrame([dict(name="R", rank=1, team="AAA", week=5, opp_txt="BBB", kick_txt="Sun 10/11 1p", spr_tot="-3 / 44", inj="",
                              n_prev=4, pg_touch=20.0, xtd=0.8, tt=27.0, note="", p_td=0.55, fair="-122")])


def test_workbook_has_top_plays_then_three_sheets_with_line_edge_formulas(tmp_path):
    elite = pd.DataFrame([dict(name="P", team="AAA", opp_txt="@BBB", kick_txt="Sun 10/4 1p", spr_tot="-3 / 44", inj="",
                               rank=1, games=3, total=900, ypg=300.0, l3=310.0, opp_rank=2, opp_alw=250.0,
                               vs_lg=0.1, proj=270.0, t_exp=34.5, t_cur=33.0)])
    empty = pd.DataFrame()
    res = Result(2026, pd.DataFrame(), {}, {"QB": elite, "RB": empty, "WR": empty}, {"QB": empty, "RB": empty, "WR": empty},
                 {"generated": pd.Timestamp("2026-10-04 12:00", tz="America/New_York"), "weeks": [4], "games": 1,
                  "data_through_week": 3, "injury_reports": {4: (4, False)}, "top_n": 10, "weak_n": 10},
                 td={"RB": td_rows(), "WR": empty, "TE": empty})
    out = tmp_path / "x.xlsx"
    write_workbook(res, out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Top plays", "1 Elite vs Weak D", "2 Backups", "3 Anytime TD"]
    ws = wb["1 Elite vs Weak D"]
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=IF(")]
    assert len(formulas) == 1 and "ISNUMBER" in formulas[0]
    assert ws["A7"].value == "P"
    ws3 = wb["3 Anytime TD"]
    f3 = [c for row in ws3.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=IF(")]
    assert len(f3) == 1 and "ISNUMBER" in f3[0].value and "100/(" in f3[0].value        # odds -> implied probability, edge = P(TD) - implied
    assert "%" in f3[0].number_format
    assert ws3["A7"].value == "R" and ws3["A6"].value == "Player"
    heads = [c.value for c in ws3[6] if c.value]
    assert "P(TD)" in heads and "Book odds" in heads and "Fair" in heads
    assert ws3["A5"].value == "RB - Anytime TD"


def test_top_plays_sheet_lists_picks_and_the_watchlist(tmp_path):
    empty = pd.DataFrame()
    res = Result(2026, pd.DataFrame(), {}, {"QB": empty, "RB": empty, "WR": empty}, {"QB": empty, "RB": empty, "WR": empty},
                 {"generated": pd.Timestamp("2026-10-04 12:00", tz="America/New_York"), "weeks": [4], "games": 1,
                  "data_through_week": 3, "injury_reports": {}, "top_n": 10, "weak_n": 10,
                  "odds": {"pulled": "2026-10-04T15:55:00Z", "book_lines": 10}},
                 td={"RB": empty, "WR": empty, "TE": empty},
                 picks=[dict(rank=1, grade="A", name="Josh Allen", team="BUF", opp_txt="@LA", kick_txt="Mon 10/12 8:15p", label="Pass yds",
                             side="Under", line=244.5, price=-112, p=0.71, implied=0.528, edge=0.181, ev=0.34,
                             reasons=["proj 210 vs line 244.5", "season 260/g (hot start, regresses)"])],
                 watch=[dict(kind="yards", name="Tee Higgins", team="CIN", opp_txt="@MIA", kick_txt="Sun 10/11 1p", label="Rec yds",
                             fair_line=58.5, over_at=54.5, under_at=61.5, reasons=["proj 67"]),
                        dict(kind="td", name="James Cook", team="BUF", opp_txt="@LA", kick_txt="Mon", label="Anytime TD", fair="-144",
                             yes_at=-125, reasons=["0.73 exp TD/g"])])
    out = tmp_path / "x.xlsx"
    write_workbook(res, out)
    ws = load_workbook(out)["Top plays"]
    vals = [[c.value for c in row] for row in ws.iter_rows()]
    flat = [v for row in vals for v in row if v is not None]
    assert ws["A1"].value == "Top plays: the model against DraftKings"
    assert "Josh Allen" in flat and "Under" in flat and 244.5 in flat and -112 in flat
    assert "proj 210 vs line 244.5; season 260/g (hot start, regresses)" in flat
    assert any(isinstance(v, str) and v.startswith("Watchlist") for v in flat)
    assert "Tee Higgins" in flat and 54.5 in flat and 61.5 in flat and "James Cook" in flat and -125 in flat
